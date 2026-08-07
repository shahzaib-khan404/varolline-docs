"""
Unit tests for all PDF processing services (Milestones 6–15).

Each test is self-contained:
  - Builds minimal test inputs in-memory (no disk I/O, no network).
  - Asserts structural correctness of the output (page count, file magic, content).
  - Does NOT require a running database, Redis, or Celery worker.
"""

from __future__ import annotations

import io
import zipfile

import fitz
import pytest
from docx import Document as DocxDocument
from PIL import Image


# ── shared fixtures ───────────────────────────────────────────────────────────


def _make_pdf(pages: int = 2, text_prefix: str = "Page") -> bytes:
    """Return a minimal in-memory PDF with *pages* text pages."""
    doc = fitz.open()
    for i in range(pages):
        page = doc.new_page(width=595, height=842)
        # Keep text long enough (>20 chars) for the OCR native-text threshold
        page.insert_text((72, 100), f"{text_prefix} number {i + 1} — test content here", fontsize=14)
    result = doc.tobytes()
    doc.close()
    return result


def _make_docx(paragraphs: list[str] | None = None) -> bytes:
    """Return a minimal in-memory DOCX."""
    doc = DocxDocument()
    for text in (paragraphs or ["Hello from Word.", "Second paragraph."]):
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_png(width: int = 100, height: int = 80, color=(200, 100, 50)) -> bytes:
    img = Image.new("RGB", (width, height), color)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _page_count(pdf_bytes: bytes) -> int:
    with fitz.open(stream=pdf_bytes, filetype="pdf") as d:
        return d.page_count


def _assert_pdf(data: bytes) -> None:
    assert data[:4] == b"%PDF", "Output is not a valid PDF"


# ── Milestone 6: Merge PDF ────────────────────────────────────────────────────


class TestMergePdf:
    def test_merges_two_pdfs(self):
        from app.services.pdf.merger import merge_pdfs

        a = _make_pdf(2)
        b = _make_pdf(3)
        result = merge_pdfs([a, b])
        _assert_pdf(result)
        assert _page_count(result) == 5

    def test_single_pdf_passthrough(self):
        from app.services.pdf.merger import merge_pdfs

        src = _make_pdf(4)
        result = merge_pdfs([src])
        _assert_pdf(result)
        assert _page_count(result) == 4

    def test_empty_list_raises(self):
        from app.services.pdf.merger import merge_pdfs

        with pytest.raises(Exception):
            merge_pdfs([])


# ── Milestone 7: Split PDF ────────────────────────────────────────────────────


class TestSplitPdf:
    def test_split_all_pages(self):
        from app.services.pdf.splitter import split_pdf

        src = _make_pdf(4)
        result = split_pdf(src, mode="all", base_name="test")
        zf = zipfile.ZipFile(io.BytesIO(result))
        assert len(zf.namelist()) == 4

    def test_split_ranges(self):
        from app.services.pdf.splitter import split_pdf

        src = _make_pdf(6)
        result = split_pdf(src, mode="ranges", ranges="1-3,5", base_name="test")
        zf = zipfile.ZipFile(io.BytesIO(result))
        assert len(zf.namelist()) == 2  # two segments

    def test_split_every_n(self):
        from app.services.pdf.splitter import split_pdf

        src = _make_pdf(6)
        result = split_pdf(src, mode="every_n", n=2, base_name="test")
        zf = zipfile.ZipFile(io.BytesIO(result))
        assert len(zf.namelist()) == 3  # 6 / 2 = 3 chunks


# ── Milestone 8: Compress PDF ─────────────────────────────────────────────────


class TestCompressPdf:
    def test_compress_returns_pdf(self):
        from app.services.pdf.compressor import compress_pdf

        src = _make_pdf(2)
        result = compress_pdf(src, quality="ebook")
        _assert_pdf(result)
        assert _page_count(result) == 2

    def test_all_quality_presets_succeed(self):
        from app.services.pdf.compressor import compress_pdf

        src = _make_pdf(1)
        for preset in ("screen", "ebook", "printer"):
            out = compress_pdf(src, quality=preset)
            _assert_pdf(out)

    def test_invalid_preset_raises(self):
        from app.services.pdf.compressor import compress_pdf

        with pytest.raises(Exception):
            compress_pdf(_make_pdf(1), quality="ultra")


# ── Milestone 9a: PDF to Image ────────────────────────────────────────────────


class TestPdfToImage:
    def test_png_output_is_zip(self):
        from app.services.pdf.pdf_to_image import pdf_to_images

        src = _make_pdf(3)
        result = pdf_to_images(src, fmt="png", dpi=72)
        zf = zipfile.ZipFile(io.BytesIO(result))
        assert len(zf.namelist()) == 3
        assert all(n.endswith(".png") for n in zf.namelist())

    def test_jpg_output(self):
        from app.services.pdf.pdf_to_image import pdf_to_images

        src = _make_pdf(2)
        result = pdf_to_images(src, fmt="jpg", dpi=72)
        zf = zipfile.ZipFile(io.BytesIO(result))
        assert all(n.endswith(".jpg") for n in zf.namelist())

    def test_dpi_affects_image_size(self):
        from app.services.pdf.pdf_to_image import pdf_to_images

        src = _make_pdf(1)
        lo = zipfile.ZipFile(io.BytesIO(pdf_to_images(src, fmt="png", dpi=72)))
        hi = zipfile.ZipFile(io.BytesIO(pdf_to_images(src, fmt="png", dpi=150)))
        lo_size = lo.read(lo.namelist()[0])
        hi_size = hi.read(hi.namelist()[0])
        assert len(hi_size) > len(lo_size)


# ── Milestone 9b: Image to PDF ────────────────────────────────────────────────


class TestImageToPdf:
    def test_single_image_to_pdf(self):
        from app.services.pdf.image_to_pdf import images_to_pdf

        png = _make_png()
        result = images_to_pdf([png])
        _assert_pdf(result)
        assert _page_count(result) == 1

    def test_multiple_images_become_multiple_pages(self):
        from app.services.pdf.image_to_pdf import images_to_pdf

        images = [_make_png(color=(c, 50, 50)) for c in (100, 150, 200)]
        result = images_to_pdf(images)
        _assert_pdf(result)
        assert _page_count(result) == 3


# ── Milestone 10a: PDF to Word ────────────────────────────────────────────────


class TestPdfToWord:
    def test_returns_docx_bytes(self):
        from app.services.pdf.pdf_to_word import pdf_to_word

        src = _make_pdf(2, text_prefix="Test content page")
        result = pdf_to_word(src)
        # DOCX is a ZIP — magic bytes are "PK"
        assert result[:2] == b"PK", "Output is not a DOCX (ZIP) file"

    def test_extracted_text_present(self):
        from app.services.pdf.pdf_to_word import pdf_to_word

        src = _make_pdf(1, text_prefix="UniqueMarker")
        docx_bytes = pdf_to_word(src)
        doc = DocxDocument(io.BytesIO(docx_bytes))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "UniqueMarker" in full_text

    def test_multipage_pdf_has_content(self):
        from app.services.pdf.pdf_to_word import pdf_to_word

        src = _make_pdf(3)
        result = pdf_to_word(src)
        doc = DocxDocument(io.BytesIO(result))
        assert len(doc.paragraphs) >= 3


# ── Milestone 10b: Word to PDF ────────────────────────────────────────────────


class TestWordToPdf:
    def test_returns_valid_pdf(self):
        from app.services.pdf.word_to_pdf import word_to_pdf

        docx = _make_docx()
        result = word_to_pdf(docx)
        _assert_pdf(result)

    def test_headed_docx_converts(self):
        from app.services.pdf.word_to_pdf import word_to_pdf

        doc = DocxDocument()
        doc.add_heading("Main Title", level=1)
        doc.add_heading("Subtitle", level=2)
        doc.add_paragraph("Body paragraph with content.")
        buf = io.BytesIO()
        doc.save(buf)
        result = word_to_pdf(buf.getvalue())
        _assert_pdf(result)

    def test_table_docx_converts(self):
        from app.services.pdf.word_to_pdf import word_to_pdf

        doc = DocxDocument()
        tbl = doc.add_table(rows=3, cols=3)
        for i, row in enumerate(tbl.rows):
            for j, cell in enumerate(row.cells):
                cell.text = f"R{i}C{j}"
        buf = io.BytesIO()
        doc.save(buf)
        result = word_to_pdf(buf.getvalue())
        _assert_pdf(result)


# ── Milestone 11: Rotate PDF ──────────────────────────────────────────────────


class TestRotatePdf:
    def test_rotate_all_90(self):
        from app.services.pdf.rotate_pdf import rotate_pdf

        src = _make_pdf(3)
        result = rotate_pdf(src, angle=90, pages="all")
        _assert_pdf(result)
        with fitz.open(stream=result, filetype="pdf") as d:
            for page in d:
                assert page.rotation == 90

    def test_rotate_odd_pages(self):
        from app.services.pdf.rotate_pdf import rotate_pdf

        src = _make_pdf(4)
        result = rotate_pdf(src, angle=180, pages="odd")
        with fitz.open(stream=result, filetype="pdf") as d:
            assert d[0].rotation == 180
            assert d[1].rotation == 0
            assert d[2].rotation == 180
            assert d[3].rotation == 0

    def test_rotate_custom_range(self):
        from app.services.pdf.rotate_pdf import rotate_pdf

        src = _make_pdf(5)
        result = rotate_pdf(src, angle=270, pages="2-4")
        with fitz.open(stream=result, filetype="pdf") as d:
            assert d[0].rotation == 0
            assert d[1].rotation == 270
            assert d[3].rotation == 270
            assert d[4].rotation == 0

    def test_invalid_angle_raises(self):
        from app.services.pdf.rotate_pdf import rotate_pdf

        with pytest.raises(ValueError):
            rotate_pdf(_make_pdf(1), angle=45, pages="all")


# ── Milestone 12: Delete Pages ────────────────────────────────────────────────


class TestDeletePages:
    def test_delete_single_page(self):
        from app.services.pdf.delete_pages import delete_pages

        src = _make_pdf(4)
        result = delete_pages(src, pages="2")
        _assert_pdf(result)
        assert _page_count(result) == 3

    def test_delete_range(self):
        from app.services.pdf.delete_pages import delete_pages

        src = _make_pdf(5)
        result = delete_pages(src, pages="2-4")
        assert _page_count(result) == 2

    def test_delete_even_pages(self):
        from app.services.pdf.delete_pages import delete_pages

        src = _make_pdf(6)
        result = delete_pages(src, pages="even")
        assert _page_count(result) == 3

    def test_cannot_delete_all_pages(self):
        from app.services.pdf.delete_pages import delete_pages

        src = _make_pdf(3)
        with pytest.raises(ValueError, match="at least one page"):
            delete_pages(src, pages="1-3")

    def test_out_of_range_pages_ignored(self):
        from app.services.pdf.delete_pages import delete_pages

        src = _make_pdf(3)
        # Page 99 doesn't exist — should not raise, just skip
        result = delete_pages(src, pages="1,99")
        assert _page_count(result) == 2


# ── Milestone 13: Watermark PDF ──────────────────────────────────────────────


class TestWatermarkPdf:
    def test_watermark_all_pages(self):
        from app.services.pdf.watermark_pdf import watermark_pdf

        src = _make_pdf(3)
        result = watermark_pdf(src, text="CONFIDENTIAL", pages="all")
        _assert_pdf(result)
        assert _page_count(result) == 3

    def test_watermark_preserves_page_count(self):
        from app.services.pdf.watermark_pdf import watermark_pdf

        src = _make_pdf(4)
        result = watermark_pdf(src, text="DRAFT", angle=0, color="#dc2626",
                               opacity=0.5, pages="odd")
        assert _page_count(result) == 4

    def test_empty_text_raises(self):
        from app.services.pdf.watermark_pdf import watermark_pdf

        with pytest.raises(ValueError, match="empty"):
            watermark_pdf(_make_pdf(1), text="   ", pages="all")

    def test_colour_parsing(self):
        from app.services.pdf.watermark_pdf import _parse_hex_color

        assert _parse_hex_color("#ff0000") == (255, 0, 0)
        assert _parse_hex_color("#0f0") == (0, 255, 0)
        assert _parse_hex_color("#1a2b3c") == (26, 43, 60)

    def test_invalid_colour_raises(self):
        from app.services.pdf.watermark_pdf import _parse_hex_color

        with pytest.raises(ValueError):
            _parse_hex_color("notacolor")


# ── Milestone 14: Password Protect ───────────────────────────────────────────


class TestPasswordProtect:
    def test_encrypted_needs_password(self):
        from app.services.pdf.password_protect import password_protect

        src = _make_pdf(2)
        result = password_protect(src, user_password="secret123")
        _assert_pdf(result)
        with fitz.open(stream=result, filetype="pdf") as d:
            assert d.needs_pass

    def test_correct_password_authenticates(self):
        from app.services.pdf.password_protect import password_protect

        src = _make_pdf(1)
        result = password_protect(src, user_password="mypass!")
        with fitz.open(stream=result, filetype="pdf") as d:
            assert d.authenticate("mypass!")

    def test_wrong_password_fails(self):
        from app.services.pdf.password_protect import password_protect

        src = _make_pdf(1)
        result = password_protect(src, user_password="correct")
        with fitz.open(stream=result, filetype="pdf") as d:
            assert not d.authenticate("wrong")

    def test_owner_password_accepted(self):
        from app.services.pdf.password_protect import password_protect

        src = _make_pdf(1)
        result = password_protect(src, user_password="user", owner_password="owner")
        with fitz.open(stream=result, filetype="pdf") as d:
            assert d.authenticate("user")

    def test_empty_password_raises(self):
        from app.services.pdf.password_protect import password_protect

        with pytest.raises(ValueError, match="empty"):
            password_protect(_make_pdf(1), user_password="  ")


# ── Milestone 15: OCR ────────────────────────────────────────────────────────


class TestOcrPdf:
    def test_native_text_extraction_txt(self):
        from app.services.pdf.ocr_pdf import ocr_pdf

        src = _make_pdf(2, text_prefix="NativeText")
        result = ocr_pdf(src, output_format="txt", language="eng")
        text = result.decode("utf-8")
        assert "NativeText" in text
        assert "--- Page 1 ---" in text
        assert "--- Page 2 ---" in text

    def test_native_text_extraction_pdf(self):
        from app.services.pdf.ocr_pdf import ocr_pdf

        src = _make_pdf(2)
        result = ocr_pdf(src, output_format="pdf", language="eng")
        _assert_pdf(result)
        assert _page_count(result) == 2

    def test_invalid_output_format_raises(self):
        from app.services.pdf.ocr_pdf import ocr_pdf

        with pytest.raises(ValueError, match="Unsupported output format"):
            ocr_pdf(_make_pdf(1), output_format="xlsx")

    def test_page_set_parser(self):
        from app.services.pdf.ocr_pdf import _parse_page_set

        assert _parse_page_set("all", 4) == [0, 1, 2, 3]
        assert _parse_page_set("odd", 4) == [0, 2]
        assert _parse_page_set("even", 4) == [1, 3]
        assert _parse_page_set("1,3", 4) == [0, 2]
        assert _parse_page_set("2-4", 5) == [1, 2, 3]

    def test_scanned_page_raises_without_tesseract(self, monkeypatch):
        """Blank pages trigger Tesseract path — should raise RuntimeError when absent."""
        from app.services.pdf.ocr_pdf import ocr_pdf
        import app.services.pdf.ocr_pdf as ocr_module
        import shutil

        # Force Tesseract lookup to always return None
        monkeypatch.setattr(shutil, "which", lambda _: None)
        monkeypatch.setattr(ocr_module, "_file_exists", lambda _: False)

        blank_doc = fitz.open()
        blank_doc.new_page()
        blank_pdf = blank_doc.tobytes()
        blank_doc.close()

        with pytest.raises(RuntimeError, match="brew install tesseract"):
            ocr_pdf(blank_pdf, output_format="txt", language="eng")


# ── Security: File Validator ──────────────────────────────────────────────────


class TestFileValidator:
    def test_clean_filename_unchanged(self):
        from app.core.file_validator import sanitize_filename

        assert sanitize_filename("invoice.pdf") == "invoice.pdf"

    def test_path_traversal_stripped(self):
        from app.core.file_validator import sanitize_filename

        result = sanitize_filename("../../etc/passwd")
        assert ".." not in result
        assert "/" not in result

    def test_empty_name_becomes_file(self):
        from app.core.file_validator import sanitize_filename

        assert sanitize_filename("") == "file"

    def test_long_stem_clamped(self):
        from app.core.file_validator import sanitize_filename

        long_name = "a" * 300 + ".pdf"
        result = sanitize_filename(long_name)
        stem = result.rsplit(".", 1)[0]
        assert len(stem) <= 120

    def test_shell_specials_replaced(self):
        from app.core.file_validator import sanitize_filename

        result = sanitize_filename("evil<script>alert(1)</script>.pdf")
        assert "<" not in result
        assert ">" not in result

    def test_unicode_preserved(self):
        from app.core.file_validator import sanitize_filename

        result = sanitize_filename("résumé.pdf")
        assert "résumé" in result

    def test_null_byte_removed(self):
        from app.core.file_validator import sanitize_filename

        result = sanitize_filename("file\x00name.pdf")
        assert "\x00" not in result
